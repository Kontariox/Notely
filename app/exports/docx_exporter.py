import io
from typing import Any, Dict, List, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor


def create_lesson_docx(
    title: str,
    subject: str,
    lesson_date: Optional[str],
    notes_markdown: str,
    events: Optional[List[Dict[str, Any]]] = None,
    transcription: Optional[str] = None
) -> io.BytesIO:
    """
    Generates a professionally formatted Word DOCX file from lesson notes and metadata.
    Returns in-memory BytesIO buffer ready for HTTP streaming/download.
    """
    doc = Document()

    # Set document title
    doc_title = doc.add_heading(title, level=0)
    doc_title.runs[0].font.size = Pt(22)
    doc_title.runs[0].font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Subtitle with metadata
    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(f"Przedmiot: {subject or 'Nieokreślony'}   |   Data: {lesson_date or 'Brak daty'}")
    meta_run.font.size = Pt(10)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph()  # spacing

    # Events summary section if present
    if events:
        doc.add_heading("Wykryte wydarzenia i terminy", level=1)
        for ev in events:
            ev_p = doc.add_paragraph(style="List Bullet")
            t_run = ev_p.add_run(f"{ev.get('title', 'Wydarzenie')}: ")
            t_run.bold = True
            date_info = ev.get('date') or ev.get('raw_date_expression') or 'Brak określonej daty'
            ev_p.add_run(f"{date_info} ({ev.get('type')}) - {ev.get('description', '')}")

    # Process notes markdown into DOCX headings and paragraphs
    if notes_markdown:
        lines = notes_markdown.splitlines()
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            if line_str.startswith("# "):
                doc.add_heading(line_str[2:].strip(), level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:].strip(), level=2)
            elif line_str.startswith("### "):
                doc.add_heading(line_str[4:].strip(), level=3)
            elif line_str.startswith("- ") or line_str.startswith("* "):
                bullet_p = doc.add_paragraph(style="List Bullet")
                bullet_p.add_run(line_str[2:].strip())
            else:
                p = doc.add_paragraph()
                p.add_run(line_str)

    # Optional transcription appendix
    if transcription:
        doc.add_page_break()
        doc.add_heading("Pełna transkrypcja nagrania", level=1)
        doc.add_paragraph(transcription)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
